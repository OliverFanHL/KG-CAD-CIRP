"""Create editable journal manuscripts with LaTeX source and three-rule tables.

Reads the revision-8 LaTeX prose and frozen comparison tables. Figures are
high-resolution source images; formulas are editable LaTeX text, not rendered math.
"""
from pathlib import Path
from copy import deepcopy
import csv,hashlib,json,re
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
TITLE='Knowledge-guided relational evidence for functional CAD retrieval'
M='http://schemas.openxmlformats.org/officeDocument/2006/math'
def node(tag,text=None,**attrs):
    e=etree.Element('{'+M+'}'+tag)
    for k,v in attrs.items():e.set('{'+M+'}'+k,str(v))
    if text is not None:e.text=str(text)
    return e
def mr(s,plain=False,bold=False):
    e=node('r');pr=node('rPr')
    if plain:pr.append(node('sty',val='p'))
    elif bold:pr.append(node('sty',val='bi'))
    if len(pr):e.append(pr)
    t=node('t',s);t.set(qn('xml:space'),'preserve');e.append(t)
    return e
def container(tag,items):
    e=node(tag)
    if isinstance(items,str):items=[mr(items)]
    for a in items:e.append(a)
    return e
def fraction(a,b):
    e=node('f');e.append(container('num',a));e.append(container('den',b));return e
def script(base,sub=None,sup=None):
    tag='sSubSup' if sub is not None and sup is not None else 'sSub' if sub is not None else 'sSup'
    e=node(tag);e.append(container('e',[base]))
    if sub is not None:e.append(container('sub',sub))
    if sup is not None:e.append(container('sup',sup))
    return e

SYMBOLS={'rho':'ρ','eta':'η','in':'∈','ldots':'…','times':'×','pm':'±','leq':'≤',
    'sum':'∑','mathsf':'','qquad':'  ','quad':' ','cos':'cos','max':'max','min':'min'}
class Math:
    def __init__(self,s):self.s=s;self.i=0
    def group(self):
        if self.i<len(self.s) and self.s[self.i]=='{':
            self.i+=1;v=self.parse('}');return v
        return [self.atom()]
    def atom(self):
        s=self.s
        if self.i>=len(s):return mr('')
        c=s[self.i];self.i+=1
        if c=='{':
            return container('box',[container('e',self.parse('}'))])
        if c=='\\':
            m=re.match(r'[A-Za-z]+',s[self.i:])
            if not m:
                c=s[self.i];self.i+=1
                return mr({'{':'{','}':'}',',':' ', '!':'',';':' '}.get(c,c),plain=True)
            cmd=m.group();self.i+=len(cmd)
            if cmd in ('frac','tfrac'):
                a=self.group();b=self.group();return fraction(a,b)
            if cmd in ('boldsymbol','mathrm','mathsf','text','operatorname','mbox'):
                while self.i<len(s) and s[self.i]==' ':self.i+=1
                val=self.group()
                if cmd=='boldsymbol':
                    for r in val:
                        if r.tag=='{'+M+'}r':r.insert(0,node('rPr'));r[0].append(node('sty',val='bi'))
                else:
                    for r in val:
                        if r.tag=='{'+M+'}r':r.insert(0,node('rPr'));r[0].append(node('sty',val='p'))
                return container('box',[container('e',val)])
            if cmd in ('left','right'):return self.atom()
            return mr(SYMBOLS.get(cmd,cmd),plain=cmd in ('cos','max','min'))
        return mr(c,plain=c.isdigit() or c in '+-=(),.;{}[] ')
    def parse(self,stop=None):
        vals=[]
        while self.i<len(self.s):
            if self.s[self.i]==stop:self.i+=1;break
            if self.s[self.i].isspace():self.i+=1;continue
            base=self.atom();sub=sup=None
            while self.i<len(self.s) and self.s[self.i] in '_^':
                typ=self.s[self.i];self.i+=1;v=self.group()
                if typ=='_':sub=v
                else:sup=v
            vals.append(script(base,sub,sup) if sub is not None or sup is not None else base)
        return vals

def omath(s):return container('oMath',Math(s).parse())

def clean(s):
    s=s.replace("Horv{\\'a}th",'Horváth').replace('G{\\"o}tling','Götling').replace('B{\\"u}ttcher','Büttcher')
    s=s.replace('J{\\"a}rvelin','Järvelin').replace('Kek{\\"a}l{\\"a}inen','Kekäläinen')
    s=re.sub(r'\\(?:texttt|textit|textbf)\{([^{}]*)\}',r'\1',s)
    s=s.replace('\\Full{}','Full').replace('\\Full','Full').replace('\\Proposed','KG–LLM')
    s=s.replace('\\%','%').replace('\\_','_').replace('\\&','&').replace('\\$','$')
    s=s.replace('~',' ').replace('---','—').replace('--','–').replace('\\,',' ')
    s=s.replace('{','').replace('}','')
    return re.sub(r'\s+',' ',s).strip()

def field(p,instruction,value,bookmark=None):
    if bookmark:
        start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bookmark[0]));start.set(qn('w:name'),bookmark[1]);p._p.append(start)
    r=OxmlElement('w:r');begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r.append(begin);p._p.append(r)
    r=OxmlElement('w:r');t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=' '+instruction+' ';r.append(t);p._p.append(r)
    r=OxmlElement('w:r');sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r.append(sep);p._p.append(r)
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=str(value);r.append(t);p._p.append(r)
    r=OxmlElement('w:r');end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r.append(end);p._p.append(r)
    if bookmark:
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bookmark[0]));p._p.append(end)

def native_equation(key):
    equations={
        'eq:geometry':r'G_i=\frac12\{1+\cos(\boldsymbol z_i,\boldsymbol z_r)\}',
        'eq:agreement':r'A_i=\boldsymbol p_i^{\mathsf T}\boldsymbol p_r',
        'eq:function':r'F_i(f)=\sum_c p_i(c)\rho_{cf}\eta_{cf}',
        'eq:base':r's_i=\frac{G_i+A_i+F_i}{3}',
        'eq:extent':r'L_i=\max_{a\in\{x,y,z\}}(x_{ia}^{\max}-x_{ia}^{\min})\leq b',
        'eq:refined':r'S_i^{\mathrm{KG-LLM}}=\frac{G_i+A_i+F_i+H_i(f)}4=\frac{3s_i+H_i(f)}4'}
    if key!='eq:coverage':return omath(equations[key])
    e=node('oMath');e.extend(Math('v(o)=').parse())
    d=node('d');pr=node('dPr');pr.append(node('begChr',val='{'));pr.append(node('endChr',val=''));d.append(pr)
    arr=node('eqArr')
    for x in ('1,  o = yes','0,  o = no','1/2,  o = unknown'):
        arr.append(container('e',[mr(x,plain=True)]))
    d.append(container('e',[arr]));e.append(d)
    e.extend(Math(r'\qquad H_i(f)=\frac14\sum_{k=1}^{4}v(o_{ifk})').parse())
    return e

class Builder:
    def __init__(self,supp=False):
        self.supp=supp;self.doc=Document();self.ids={};self.counter=100;self.eq=0;self.fig=0;self.tab=0
        self.source=(ROOT/('supplement.tex' if supp else 'main.tex')).read_text(encoding='utf8')
        self.references=re.findall(r'\\bibitem\{([^}]+)\}\s*(.*?)(?=\\bibitem|\\end\{thebibliography\})',
            (ROOT/'references.bbl').read_text(encoding='utf8'),re.S)
        self.cites={k:i for i,(k,v) in enumerate(self.references,1)}
        self.configure();self.make_ids()
    def configure(self):
        d=self.doc;s=d.sections[0]
        s.page_width=Inches(8.2677);s.page_height=Inches(11.6929)
        s.top_margin=Inches(.67);s.bottom_margin=Inches(.67);s.left_margin=Inches(.67);s.right_margin=Inches(.67)
        self.width=6.9277
        for n in ('Normal','Body Text','Caption','Title','Heading 1','Heading 2','Heading 3'):
            st=d.styles[n];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
            st.paragraph_format.space_after=Pt(5)
        st=d.styles['Normal'];st.font.size=Pt(10.5);st.paragraph_format.line_spacing=1.1
        st.font.bold=False;st.font.italic=False
        st.paragraph_format.widow_control=True
        d.styles['Title'].font.size=Pt(18);d.styles['Title'].font.bold=True
        for n,size in [('Heading 1',12),('Heading 2',11)]:
            st=d.styles[n];st.font.size=Pt(size);st.font.bold=True
            st.paragraph_format.space_before=Pt(11);st.paragraph_format.space_after=Pt(5);st.paragraph_format.keep_with_next=True
        d.styles['Caption'].font.size=Pt(9);d.styles['Caption'].font.italic=False
        d.styles['Caption'].font.bold=False
        d.styles['Caption'].paragraph_format.line_spacing=1.05
        if 'LaTeX Code' not in d.styles:
            from docx.enum.style import WD_STYLE_TYPE
            d.styles.add_style('LaTeX Code',WD_STYLE_TYPE.PARAGRAPH)
        code=d.styles['LaTeX Code'];code.font.name='Consolas';code.font.size=Pt(9)
        code.paragraph_format.line_spacing=1.0;code.paragraph_format.keep_together=True
        code.paragraph_format.space_before=Pt(5);code.paragraph_format.space_after=Pt(7)
        footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        field(footer,'PAGE','1')
        d.core_properties.title=TITLE if not self.supp else 'Supplementary methods and results '+TITLE
        d.core_properties.author='';d.core_properties.subject='Functional CAD retrieval'
        d.core_properties.keywords='CAD; knowledge graph; functional retrieval; visual language models'
        d.core_properties.comments=''
        update=OxmlElement('w:updateStyles');update.set(qn('w:val'),'false');d.settings.element.append(update)
    def bookmark(self,key):
        self.counter+=1;return self.counter,re.sub('[^A-Za-z0-9_]','_',key)
    def make_ids(self):
        for env,token in [('figure','Figure'),('table','Table')]:
            for i,m in enumerate(re.finditer(r'\\begin\{'+env+r'\*?\}.*?\\end\{'+env+r'\*?\}',self.source,re.S),1):
                key=re.search(r'\\label\{([^}]+)\}',m.group()).group(1)
                self.ids[key]=(self.bookmark(key),('S' if self.supp else '')+str(i))
        for i,key in enumerate(re.findall(r'\\label\{(eq:[^}]+)\}',self.source),1):self.ids[key]=(self.bookmark(key),str(i))
    def inline(self,p,s):
        pattern=r'(\$[^$]+\$|\\cite\{[^}]+\}|\\ref\{[^}]+\}|\\(?:texttt|textit|textbf|url)\{[^}]+\}|\\href\{[^}]+\}\{[^}]+\})'
        for bit in re.split(pattern,s):
            if not bit:continue
            if bit.startswith('$'):
                raw=bit[1:-1]
                numeric=re.fullmatch(r'[0-9.,+\-\[\] ]+(?:\\!?|\\pm|\\,|[0-9.,+\-\[\] ])*',raw)
                if raw in (r'\pm',r'\times'):
                    p.add_run({r'\pm':'±',r'\times':'×'}[raw])
                elif getattr(self,'in_table',False) or numeric:
                    plain=raw.replace(r'\!', '').replace(r'\pm',' ± ').replace(r'\,',' ').replace(r'\times','×')
                    r=p.add_run(plain);r.italic=bool(re.fullmatch(r'[A-Za-z+, ]+',plain))
                elif re.fullmatch(r'[A-Za-z](?:[,+][A-Za-z])*',raw):
                    r=p.add_run(raw);r.italic=True
                else:
                    r=p.add_run('$'+raw+'$');r.font.name='Consolas';r.font.size=Pt(9)
                continue
            if bit.startswith('\\cite{'):
                keys=bit[6:-1].split(',');p.add_run('['+', '.join(str(self.cites[k]) for k in keys)+']');continue
            if bit.startswith('\\ref{'):
                key=bit[5:-1];book,num=self.ids[key];field(p,'REF '+book[1]+' \\h',num);continue
            if bit.startswith('\\url{'):p.add_run(bit[5:-1]);continue
            if bit.startswith('\\href{'):
                url,label=re.match(r'\\href\{([^}]+)\}\{([^}]+)\}',bit).groups()
                p.add_run(clean(label)+' ('+url+')');continue
            if bit.startswith(('\\texttt{','\\textit{','\\textbf{')):
                cmd,txt=re.match(r'\\(\w+)\{(.*)\}',bit,re.S).groups();r=p.add_run(clean(txt))
                r.italic=cmd=='textit';r.bold=cmd=='textbf'
                if cmd=='texttt':r.font.name='Consolas';r.font.size=Pt(9)
                continue
            leading=' ' if bit.startswith((' ','\n','~')) else ''
            trailing=' ' if bit.endswith((' ','\n','~')) else ''
            p.add_run(leading+clean(bit)+trailing)
    def para(self,s):
        if not clean(s):return
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        self.inline(p,s);return p
    def caption(self,kind,body,key):
        self.fig+=kind=='Figure';self.tab+=kind=='Table'
        n=self.fig if kind=='Figure' else self.tab
        p=self.doc.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=False
        p.paragraph_format.keep_together=True
        p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(9)
        p.add_run(kind+' ')
        book,num=self.ids[key]
        if self.supp:
            start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(book[0]));start.set(qn('w:name'),book[1]);p._p.append(start)
            p.add_run('S');field(p,'SEQ '+kind+' \\* ARABIC',n)
            end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(book[0]));p._p.append(end)
        else:field(p,'SEQ '+kind+' \\* ARABIC',n,book)
        p.add_run('. ');self.inline(p,body)
    def figure(self,block):
        key=re.search(r'\\label\{([^}]+)\}',block).group(1)
        name=re.search(r'\\includegraphics(?:\[[^]]*\])?\{([^}]+)\}',block).group(1)
        image=ROOT/'figures'/name
        if image.suffix=='.pdf':image=image.with_suffix('.png')
        if self.doc.paragraphs and len(self.doc.paragraphs[-1].text)<110:
            self.doc.paragraphs[-1].paragraph_format.keep_with_next=True
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(0)
        shape=p.add_run().add_picture(str(image),width=Inches(self.width))
        shape._inline.docPr.set('descr',{'fig:framework':'Knowledge graph and CAD retrieval workflow',
            'fig:performance':'Mean retrieval performance and collection SD by observer and representation',
            'fig:families':'Role-specific paired mean nDCG differences and collection SD',
            'fig:evidence':'Source CAD models with role predicates and observed geometric evidence'}[key])
        caption=block[block.index('\\caption{')+9:block.index('\n\\label')-1]
        self.caption('Figure',caption,key)
    def table_style(self,table,widths):
        table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        # Three-line academic style: top, header separator, bottom only.
        tblpr=table._tbl.tblPr
        for old in tblpr.findall(qn('w:tblBorders')):tblpr.remove(old)
        borders=OxmlElement('w:tblBorders')
        for tag in ('top','left','bottom','right','insideH','insideV'):
            edge=OxmlElement('w:'+tag)
            edge.set(qn('w:val'),'single' if tag in ('top','bottom') else 'nil')
            if tag in ('top','bottom'):
                edge.set(qn('w:sz'),'8');edge.set(qn('w:color'),'000000')
            borders.append(edge)
        tblpr.append(borders)
        for row in table.rows:
            trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            for j,cell in enumerate(row.cells):
                cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for tag in ('top','left','bottom','right'):
                    e=OxmlElement('w:'+tag);e.set(qn('w:val'),'nil');borders.append(e)
                pr.append(borders);margins=OxmlElement('w:tcMar')
                for tag,val in [('top',50),('bottom',50),('left',75),('right',75)]:
                    e=OxmlElement('w:'+tag);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
                pr.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.0
        for cell in table.rows[0].cells:
            pr=cell._tc.get_or_add_tcPr()
            edge=pr.find(qn('w:tcBorders')).find(qn('w:bottom'))
            edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'000000')
            for old in pr.findall(qn('w:shd')):pr.remove(old)
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for i,row in enumerate(table.rows):
            for j,cell in enumerate(row.cells):
                cell_borders=cell._tc.get_or_add_tcPr().find(qn('w:tcBorders'))
                for name,enabled in [('top',i==0),('bottom',i==len(table.rows)-1)]:
                    if enabled:
                        edge=cell_borders.find(qn('w:'+name));edge.set(qn('w:val'),'single')
                        edge.set(qn('w:sz'),'8');edge.set(qn('w:color'),'000000')
                p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                if (not self.supp and len(widths)==2) or (self.supp and len(widths)==3):p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                if (len(widths)==5 and j>0) or (len(widths)==6 and j>1):p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if i>0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(9 if len(widths)<=5 else 8.2);r.bold=(i==0)
                if i>=len(table.rows)-(3 if self.supp and len(widths)==3 else 1):p.paragraph_format.keep_with_next=True
    def table(self,block):
        key=re.search(r'\\label\{([^}]+)\}',block).group(1)
        if key=='tab:contrasts':
            rows=list(csv.DictReader((ROOT/'data/paired_tests.csv').open()))
            data=[['Representation','Contrast','Δ nDCG@5','95% interval','Holm p','W/T/L']]
            for r in rows:
                data.append([r['geometry_kind'].capitalize(),'KG Luna − Full+J' if r['proposed']=='kg_luna' else 'KG Sol − KG Luna',
                    f"{float(r['delta']):+.4f}",f"[{float(r['ci95_lower']):+.4f}, {float(r['ci95_upper']):+.4f}]",
                    f"{float(r['holm_p']):.6f}",'/'.join(r[x] for x in ('family_wins','family_ties','family_losses'))])
            widths=[.83,1.40,.72,1.68,.77,.63]
        else:
            if key=='tab:paired':
                tabtext=(ROOT/'tables/component_ablation_v7.tex').read_text();data=[['Scoring configuration','Embedding\nnDCG@5','Embedding\nStrong P@5','Descriptor\nnDCG@5','Descriptor\nStrong P@5']]
                raw=tabtext[tabtext.index('\\midrule')+8:];widths=[2.15,1.19,1.19,1.19,1.19]
            else:
                raw=block[block.index('\\toprule')+8:];data=[]
                widths=[1.65,5.26] if key=='tab:design' else [1.18,.77,4.96]
            for line in raw.splitlines():
                if '&' not in line or not line.rstrip().endswith('\\\\'):continue
                if '\\multicolumn' in line:continue
                data.append([x.strip() for x in line.rstrip().removesuffix('\\\\').split('&')])
        table=self.doc.add_table(rows=len(data),cols=len(data[0]))
        self.in_table=True
        for row,vals in zip(table.rows,data):
            for cell,value in zip(row.cells,vals):self.inline(cell.paragraphs[0],value)
        self.in_table=False
        self.table_style(table,widths)
        caption=block[block.index('\\caption{')+9:block.index('\n\\label')-1]
        self.caption('Table',caption,key)
    def equations(self,block):
        if self.doc.paragraphs:self.doc.paragraphs[-1].paragraph_format.keep_with_next=True
        body=re.sub(r'\\(?:begin|end)\{(?:align|equation)\}','',block).strip()
        previous=0
        for match in re.finditer(r'\\label\{(eq:[^}]+)\}',body):
            key=match.group(1);raw=body[previous:match.start()].strip();previous=match.end()
            raw=raw.strip('\\,\n ');raw=raw.replace('&=','=')
            raw=raw.replace(r'\qquad ', r'\qquad'+'\n')
            if r'\begin{cases}' in raw:
                raw=raw.replace(r'\begin{cases}',r'\begin{cases}'+'\n')
                raw=raw.replace(r'\\',r'\\'+'\n').replace(r'\end{cases}','\n'+r'\end{cases}')
            if key=='eq:refined':raw=raw.replace('=\\frac{3s_i', '\n=\\frac{3s_i')
            self.eq+=1;p=self.doc.add_paragraph(style='LaTeX Code')
            r=p.add_run('\\[\n'+raw+'\n\\]');r.font.name='Consolas';r.font.size=Pt(9)
            p.add_run('\t(')
            p.paragraph_format.tab_stops.add_tab_stop(Inches(self.width-.40))
            field(p,'SEQ Equation \\* ARABIC',self.eq,self.ids[key][0]);p.add_run(')')
    def parse(self,s):
        envs=r'\\begin\{(?P<env>figure\*?|table\*?|align|equation)\}.*?\\end\{(?P=env)\}'
        heads=r'\\(?:section|subsection)\*?\{[^}]+\}'
        chunks=re.split('('+envs+'|'+heads+')',s,flags=re.S)
        sec=sub=0
        for x in chunks:
            if not x or x in ('figure','figure*','table','table*','align','equation'):continue
            if x.startswith('\\begin{figure'):self.figure(x);continue
            if x.startswith('\\begin{table'):self.table(x);continue
            if x.startswith(('\\begin{align','\\begin{equation')):self.equations(x);continue
            if x.startswith(('\\section','\\subsection')):
                txt=re.search(r'\{(.*)\}',x).group(1)
                if x.startswith('\\subsection'):sub+=1;prefix=f'{sec}.{sub} ';level=2
                elif '*' in x:prefix='';level=1
                else:sec+=1;sub=0;prefix=('S' if self.supp else '')+str(sec)+' ';level=1
                self.doc.add_paragraph(prefix+clean(txt),style='Heading '+str(level));continue
            # Formatting-only trailing commands have no semantic content.
            x=x.split('\\par\\vspace{6pt}')[0]
            for para in re.split(r'\n\s*\n',x):
                if para.strip().startswith(('\\end{document}','\\input{references.bbl}')):continue
                if para.strip():self.para(para)
    def build(self):
        self.doc.add_paragraph(TITLE if not self.supp else 'Supplementary methods and results',style='Title')
        if self.supp:self.doc.add_paragraph(TITLE,style='Subtitle')
        else:
            abstract=re.search(r'\\small\s*\n(Functional.*?)\n\s*\n\\par',self.source,re.S).group(1)
            p=self.doc.add_paragraph();p.add_run('Abstract. ').bold=True;self.inline(p,abstract)
            p=self.doc.add_paragraph();p.add_run('Keywords. ').bold=True
            p.add_run('Computer aided design; Knowledge based system; Design reuse; Knowledge graph')
        s=self.source[self.source.index('\\section{'):]
        if not self.supp:s=s[:s.index('\\par\\vspace{6pt}')]
        else:s=s[:s.index('\\end{document}')]
        self.parse(s)
        if not self.supp:
            self.doc.add_paragraph('References',style='Heading 1')
            for i,(key,entry) in enumerate(self.references,1):
                p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
                p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.first_line_indent=Inches(-.25)
                self.inline(p,f'[{i}] '+entry)
                for r in p.runs:r.font.size=Pt(8.5)
        # Black subtitle and source-native editability; no theme-blue content.
        self.doc.styles['Subtitle'].font.name='Times New Roman';self.doc.styles['Subtitle'].font.size=Pt(12)
        self.doc.styles['Subtitle'].font.color.rgb=RGBColor(0,0,0)
        self.doc.styles['Subtitle'].font.italic=False
        # Remove inherited template borders and theme fonts explicitly.
        for el in self.doc.styles.element.xpath('.//w:pBdr'):
            el.getparent().remove(el)
        for name in ('Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Caption'):
            style=self.doc.styles[name];pr=style.element.get_or_add_rPr()
            for old in pr.findall(qn('w:rFonts')):pr.remove(old)
            fonts=OxmlElement('w:rFonts')
            for attr in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+attr),'Times New Roman')
            pr.insert(0,fonts)
        for p in self.doc.paragraphs:
            for old in p._p.xpath('./w:pPr/w:pBdr'):old.getparent().remove(old)
            for r in p.runs:
                if p.style.name=='LaTeX Code':r.font.name='Consolas';r.font.size=Pt(9)
                elif r.font.name!='Consolas':r.font.name='Times New Roman'
                if p.style.name=='Title':r.font.size=Pt(18);r.bold=True;r.font.color.rgb=RGBColor(0,0,0)
                if p.style.name=='Caption':r.font.size=Pt(9);r.bold=False
                if p.style.name.startswith('Heading'):r.bold=True;r.font.color.rgb=RGBColor(0,0,0)
        path=ROOT/('supplement.docx' if self.supp else 'manuscript.docx');self.doc.save(path)
        return dict(file=path.name,paragraphs=len(self.doc.paragraphs),tables=len(self.doc.tables),
            figures=self.fig,numbered_latex_equations=self.eq,rendered_math_objects=0,
            table_style='Three horizontal rules; no vertical or body rules; no shading',
            references=0 if self.supp else len(self.references))

if __name__=='__main__':
    contrasts=list(csv.DictReader((ROOT/'data/paired_tests.csv').open()))
    tex=[r'\begin{tabular}{@{}llrrrr@{}}',r'\toprule',
         r'Representation & Contrast & $\Delta$ nDCG@5 & 95\% interval & Holm $p$ & W/T/L \\',r'\midrule']
    for r in contrasts:
        name='KG Luna -- Full+$J$' if r['proposed']=='kg_luna' else 'KG Sol -- KG Luna'
        tex.append(' & '.join([r['geometry_kind'].capitalize(),name,f"{float(r['delta']):+.4f}",
            f"[{float(r['ci95_lower']):+.4f}, {float(r['ci95_upper']):+.4f}]",f"{float(r['holm_p']):.6f}",
            '/'.join(r[x] for x in ('family_wins','family_ties','family_losses'))])+r' \\')
    tex.extend([r'\bottomrule',r'\end{tabular}'])
    (ROOT/'tables/paired_contrasts_v8.tex').write_text('\n'.join(tex)+'\n',encoding='utf8')
    records=[Builder(False).build(),Builder(True).build()]
    (ROOT/'provenance/word_structure.json').write_text(json.dumps(records,indent=2)+'\n')
    print(json.dumps(records,indent=2))
